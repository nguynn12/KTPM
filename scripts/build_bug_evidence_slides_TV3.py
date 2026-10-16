#!/usr/bin/env python3
"""
scripts/build_bug_evidence_slides_TV3.py
Tạo file trình chiếu PowerPoint 2 slide chuẩn cho BUG-01 và BUG-03
Nguyễn Khiêm Thuận - MSSV 2312764 - TV3 - Seed 303.
Lưu file vào: docs/slides/TV3_BUG_EVIDENCE_303.pptx
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EVIDENCE_DIR = os.path.join(ROOT_DIR, "docs", "evidence")
SLIDES_DIR = os.path.join(ROOT_DIR, "docs", "slides")
os.makedirs(SLIDES_DIR, exist_ok=True)
OUTPUT_PPTX = os.path.join(SLIDES_DIR, "TV3_BUG_EVIDENCE_303.pptx")

# Palette
DARK_BG = RGBColor(24, 24, 37)       # #181825
CARD_BG = RGBColor(30, 30, 46)       # #1E1E2E
CARD_BORDER = RGBColor(69, 71, 90)   # #45475A
TEXT_LIGHT = RGBColor(205, 214, 244) # #CDD6F4
TEXT_MUTED = RGBColor(166, 173, 200) # #A6ADC8
RED_ACCENT = RGBColor(243, 139, 168) # #F38BA8
GREEN_ACCENT = RGBColor(166, 227, 161) # #A6E3A1
BLUE_ACCENT = RGBColor(137, 180, 250) # #89B4FA
YELLOW_ACCENT = RGBColor(249, 226, 175) # #F9E2AF


def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG


def add_header(slide, title_text, badge_text, badge_color):
    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(9.5), Inches(0.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    
    p2 = tf.add_paragraph()
    p2.text = "Thành viên 3: Nguyễn Khiêm Thuận · MSSV: 2312764 · Seed: 303 | Paperless-ngx v3.1.3 | Schemathesis 4.30.1"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED

    # Badge Box (Right side)
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(0.4), Inches(2.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = badge_color
    badge.line.color.rgb = badge_color
    btf = badge.text_frame
    btf.margin_top = btf.margin_bottom = btf.margin_left = btf.margin_right = 0
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = badge_text
    bp.font.size = Pt(12)
    bp.font.bold = True
    bp.font.color.rgb = RGBColor(17, 17, 27)


def build_slide1(prs):
    # Slide 1: BUG-01
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    set_slide_background(slide)
    
    add_header(slide, "BUG-01: Dữ Liệu Permissions Sai Kiểu Gây Lỗi Máy Chủ HTTP 500", "ĐÃ TÁI HIỆN (HIGH)", RED_ACCENT)
    
    # 2 Images side by side
    img1_path = os.path.join(EVIDENCE_DIR, "TV3_BUG01_correspondents_500.png")
    img2_path = os.path.join(EVIDENCE_DIR, "TV3_BUG01_document_types_500.png")
    
    img_w = Inches(5.9)
    img_h = Inches(3.32)
    top_y = Inches(1.3)
    
    if os.path.exists(img1_path):
        slide.shapes.add_picture(img1_path, Inches(0.6), top_y, width=img_w, height=img_h)
    if os.path.exists(img2_path):
        slide.shapes.add_picture(img2_path, Inches(6.8), top_y, width=img_w, height=img_h)
        
    # Bottom Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.75), Inches(12.1), Inches(2.35))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1)
    
    ctf = card.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Inches(0.25)
    ctf.margin_right = Inches(0.25)
    ctf.margin_top = Inches(0.18)
    ctf.margin_bottom = Inches(0.15)
    
    p_card_title = ctf.paragraphs[0]
    p_card_title.text = "TÓM TẮT KỸ THUẬT & ĐỐI CHIẾU HỢP ĐỒNG API (BUG-01)"
    p_card_title.font.size = Pt(13)
    p_card_title.font.bold = True
    p_card_title.font.color.rgb = YELLOW_ACCENT
    
    bullet_items = [
        ("• Endpoint & Dữ liệu kích hoạt: ", True, BLUE_ACCENT, "POST /api/correspondents/ và POST /api/document_types/ với payload set_permissions.view = false (boolean thay vì object).", False, TEXT_LIGHT),
        ("• Thực tế hệ thống (Actual): ", True, RED_ACCENT, "HTTP 500 Internal Server Error dạng HTML; máy chủ Django gặp ngoại lệ chưa được xử lý (Unhandled Server Error).", False, TEXT_LIGHT),
        ("• Kỳ vọng hợp đồng (Expected): ", True, GREEN_ACCENT, "HTTP 400 Bad Request; theo đặc tả OpenAPI, view phải là object chứa danh sách ID người dùng/nhóm ({users: [], groups: []}).", False, TEXT_LIGHT),
        ("• Trạng thái & Kết luận: ", True, YELLOW_ACCENT, "ĐÃ TÁI HIỆN THÀNH CÔNG 2 LẦN ĐỘC LẬP trên cả hai tài nguyên. Lỗi kiến trúc dùng chung ở tầng PermissionsMixin; chưa kết luận cụ thể nguyên nhân mã nguồn do thiếu traceback server nội bộ container.", False, TEXT_LIGHT),
    ]
    
    for prefix, pre_bold, pre_color, text, body_bold, body_color in bullet_items:
        p = ctf.add_paragraph()
        p.space_after = Pt(4)
        run1 = p.add_run()
        run1.text = prefix
        run1.font.bold = pre_bold
        run1.font.size = Pt(11)
        run1.font.color.rgb = pre_color
        
        run2 = p.add_run()
        run2.text = text
        run2.font.bold = body_bold
        run2.font.size = Pt(11)
        run2.font.color.rgb = body_color

    # Speaker Notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = (
        "SPEAKER NOTES — SLIDE 1 (BUG-01):\n"
        "- Người trình bày: Nguyễn Khiêm Thuận (TV3) - MSSV: 2312764 - Seed: 303.\n"
        "- Bằng chứng gốc: docs/evidence/TV3_BUG01_correspondents_500.png và docs/evidence/TV3_BUG01_document_types_500.png.\n"
        "- Phiên bản đã xác minh: Paperless-ngx v3.1.3 (Docker Desktop 29.7.2, Windows 11).\n"
        "- Chi tiết tái hiện: Lỗi do TV2 phát hiện đầu tiên trên Tag. TV3 đã tái hiện độc lập 100% trên Correspondents và Document Types bằng 2 lượt gọi HTTP thực tế.\n"
        "- Điểm lưu ý trả lời giảng viên: Không tự suy đoán lỗi do TypeError hay ép kiểu ORM nếu chưa trích xuất traceback từ log container; ghi nhận chuẩn xác là Unhandled Server Error ở lớp PermissionsMixin dùng chung."
    )


def build_slide2(prs):
    # Slide 2: BUG-03
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    set_slide_background(slide)
    
    add_header(slide, "BUG-03: Phản Hồi Tạo Mới Correspondent Không Khớp OpenAPI Schema", "SAI LỆCH HỢP ĐỒNG (MEDIUM)", YELLOW_ACCENT)
    
    # 2 Images side by side (Response 201 vs Schema Required)
    img1_path = os.path.join(EVIDENCE_DIR, "TV3_BUG03_response_201.png")
    img2_path = os.path.join(EVIDENCE_DIR, "TV3_BUG03_schema_required.png")
    
    img_w = Inches(5.9)
    img_h = Inches(3.32)
    top_y = Inches(1.3)
    
    if os.path.exists(img1_path):
        slide.shapes.add_picture(img1_path, Inches(0.6), top_y, width=img_w, height=img_h)
    if os.path.exists(img2_path):
        slide.shapes.add_picture(img2_path, Inches(6.8), top_y, width=img_w, height=img_h)
        
    # Bottom Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.75), Inches(12.1), Inches(2.35))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1)
    
    ctf = card.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Inches(0.25)
    ctf.margin_right = Inches(0.25)
    ctf.margin_top = Inches(0.18)
    ctf.margin_bottom = Inches(0.15)
    
    p_card_title = ctf.paragraphs[0]
    p_card_title.text = "TÓM TẮT KỸ THUẬT & ĐỐI CHIẾU HỢP ĐỒNG API (BUG-03)"
    p_card_title.font.size = Pt(13)
    p_card_title.font.bold = True
    p_card_title.font.color.rgb = YELLOW_ACCENT
    
    bullet_items = [
        ("• Thao tác & Endpoint: ", True, BLUE_ACCENT, "POST /api/correspondents/ tạo mới bản ghi thành công với mã phản hồi HTTP 201 Created (ID tự sinh).", False, TEXT_LIGHT),
        ("• Trường bắt buộc bị thiếu đã xác minh: ", True, RED_ACCENT, "Phản hồi 201 thực tế thiếu chính xác 3 trường required theo schema: document_count, last_correspondence, và permissions.", False, TEXT_LIGHT),
        ("• Kiểm tra điều kiện full_perms: ", True, GREEN_ACCENT, "Trường permissions chỉ trả về khi gọi GET với full_perms=true, nhưng schema OpenAPI lại đánh required vô điều kiện trong response 201 của POST.", False, TEXT_LIGHT),
        ("• Phân loại sai lệch hợp đồng: ", True, YELLOW_ACCENT, "Sai lệch hợp đồng API (Schema Contract Violation) giữa drf-spectacular và response serializer; chưa quy kết lỗi ở backend hay tài liệu nếu chưa đủ bằng chứng. Hồ sơ TV3 chỉ khẳng định lỗi trên Correspondents.", False, TEXT_LIGHT),
    ]
    
    for prefix, pre_bold, pre_color, text, body_bold, body_color in bullet_items:
        p = ctf.add_paragraph()
        p.space_after = Pt(4)
        run1 = p.add_run()
        run1.text = prefix
        run1.font.bold = pre_bold
        run1.font.size = Pt(11)
        run1.font.color.rgb = pre_color
        
        run2 = p.add_run()
        run2.text = text
        run2.font.bold = body_bold
        run2.font.size = Pt(11)
        run2.font.color.rgb = body_color

    # Speaker Notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = (
        "SPEAKER NOTES — SLIDE 2 (BUG-03):\n"
        "- Người trình bày: Nguyễn Khiêm Thuận (TV3) - MSSV: 2312764 - Seed: 303.\n"
        "- Bằng chứng gốc: docs/evidence/TV3_BUG03_response_201.png, docs/evidence/TV3_BUG03_schema_required.png, docs/openapi_TV3.json.\n"
        "- Mã băm schema SHA256: 888b0f89ec0a9f65293f5320a86b1810d7904f93033178b567e313d951b12e3a.\n"
        "- Chi tiết tái hiện: Đã xác thực qua 2 lần POST tạo mới ID 111 và ID 112 trên instance local. Cả 2 lần response 201 đều chỉ có 8 trường và thiếu 3 trường required.\n"
        "- Điểm lưu ý trả lời giảng viên: Phân tích rõ nguyên nhân bất đồng bộ do drf-spectacular dùng chung một schema Correspondent cho cả GET/POST, trong khi trường permissions cần cờ full_perms và 2 trường count/date là thuộc tính annotate chỉ tính toán khi truy vấn danh sách."
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    print("[INFO] Bắt đầu khởi tạo 2 slide minh chứng lỗi TV3...")
    build_slide1(prs)
    build_slide2(prs)
    
    prs.save(OUTPUT_PPTX)
    print(f"[THÀNH CÔNG] Đã tạo file slide hoàn chỉnh: {OUTPUT_PPTX}")
    print(f"  - Kích thước: 13.333 x 7.5 inches (Widescreen 16:9)")
    print(f"  - Số slide: {len(prs.slides)}")
    print(f"  - Dung lượng: {os.path.getsize(OUTPUT_PPTX) / 1024:.1f} KB")


if __name__ == "__main__":
    main()
