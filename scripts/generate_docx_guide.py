#!/usr/bin/env python3
"""
scripts/generate_docx_guide.py
Tạo tài liệu Word (.docx) Hướng dẫn sử dụng và Quy trình kiểm thử hệ thống Paperless-ngx
Phục vụ đề tài Kiểm thử phần mềm - Nhóm 17 (TV4 - Nguyễn Nhất Minh - MSSV: 2312690)
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Cấu hình lề trang (1 inch = 1440 dxa)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Tiêu đề tài liệu
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = title_p.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — TRƯỜNG ĐẠI HỌC ĐÀ LẠT\nKHOA CÔNG NGHỆ THÔNG TIN\n\n")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(70, 70, 70)

    run_title = title_p.add_run("TÀI LIỆU HƯỚNG DẪN VẬN HÀNH VÀ\nQUY TRÌNH KIỂM THỬ HỆ THỐNG PAPERLESS-NGX\n")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(24, 76, 120)

    run_desc = title_p.add_run("Đề tài R08 + K05: Kiểm thử tự động API bằng Schemathesis và Kiểm thử hộp đen thực nghiệm\n")
    run_desc.font.name = "Times New Roman"
    run_desc.font.size = Pt(13)
    run_desc.font.italic = True

    doc.add_paragraph()

    # Bảng thông tin sinh viên
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Nhóm học phần:", "Nhóm 17 (Học phần Kiểm thử phần mềm)"),
        ("Thành viên thực hiện:", "Nguyễn Nhất Minh (TV4 — Phụ trách Automation & Tổng hợp số liệu)"),
        ("Mã số sinh viên (MSSV):", "2312690"),
        ("Hạt giống kiểm thử (Seed):", "Seed 404"),
        ("Phiên bản hệ thống:", "Paperless-ngx v3.1.3 (Git Commit SHA d48663e9) · Schemathesis v4.31.0"),
    ]
    for i, (k, v) in enumerate(info_data):
        row = info_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.text = k
        c1.text = v
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.name = "Times New Roman"
        c1.paragraphs[0].runs[0].font.name = "Times New Roman"
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FAFAFA")
        set_cell_margins(c0)
        set_cell_margins(c1)

    doc.add_page_break()

    # MỤC 1: HƯỚNG DẪN KHỞI CHẠY HỆ THỐNG
    h1 = doc.add_heading("1. Hướng dẫn khởi chạy và vận hành Paperless-ngx cục bộ", level=1)
    h1.runs[0].font.name = "Times New Roman"
    h1.runs[0].font.color.rgb = RGBColor(24, 76, 120)

    p = doc.add_paragraph()
    p.add_run("Hệ thống Paperless-ngx được triển khai cục bộ thông qua Docker Compose với 3 dịch vụ cơ sở: webserver (Django/DRF), database (PostgreSQL 16), và message broker (Redis 7).\n").font.name = "Times New Roman"

    doc.add_heading("1.1. Khởi động các container Docker", level=2).runs[0].font.name = "Times New Roman"
    code_p = doc.add_paragraph()
    run_code = code_p.add_run(
        "cd D:\\Kiem-thu-phan-mem\\KTPM\n"
        "docker compose -f docker-compose.env.yml up -d\n"
        "docker compose ps"
    )
    run_code.font.name = "Consolas"
    run_code.font.size = Pt(10)

    doc.add_heading("1.2. Kiểm tra trạng thái và thông tin xác thực", level=2).runs[0].font.name = "Times New Roman"
    p2 = doc.add_paragraph()
    p2.add_run("• Giao diện Web UI: ").bold = True
    p2.add_run("http://localhost:8000 (Tài khoản: admin / Mật khẩu: admin)\n")
    p2.add_run("• Tài liệu OpenAPI Schema (Swagger UI): ").bold = True
    p2.add_run("http://localhost:8000/api/schema/swagger-ui/\n")
    p2.add_run("• Lấy Auth Token qua API:\n")
    
    code_token = doc.add_paragraph()
    c_run = code_token.add_run(
        'curl -X POST http://localhost:8000/api/token/ \\\n'
        '     -H "Content-Type: application/json" \\\n'
        '     -d \'{"username":"admin","password":"admin"}\''
    )
    c_run.font.name = "Consolas"
    c_run.font.size = Pt(10)

    # MỤC 2: QUY TRÌNH 3 VÒNG KIỂM THỬ SCHEMATHESIS
    h2 = doc.add_heading("2. Quy trình kiểm thử tự động 3 vòng chuẩn hóa bằng Schemathesis", level=1)
    h2.runs[0].font.name = "Times New Roman"
    h2.runs[0].font.color.rgb = RGBColor(24, 76, 120)

    p3 = doc.add_paragraph()
    p3.add_run("Để đảm bảo tính nhất quán trên toàn nhóm, phạm vi kiểm thử được chuẩn hóa trên ").font.name = "Times New Roman"
    p3.add_run("4 nhóm tài nguyên trọng tâm: tags, correspondents, document_types, documents").bold = True
    p3.add_run(" gồm đúng 53 operations (thay vì quét toàn bộ 163 operations của cả hệ thống).\n")

    # Vòng 1
    doc.add_heading("2.1. Vòng 1: Kiểm thử chỉ đọc (GET) — Chuẩn hóa 19 Operations", level=2).runs[0].font.name = "Times New Roman"
    p_v1 = doc.add_paragraph()
    p_v1.add_run("• Mục tiêu: Quét các truy vấn GET danh sách và chi tiết trên 4 nhóm tài nguyên.\n• Cờ tham số bắt buộc: ").font.name = "Times New Roman"
    p_v1.add_run("--include-path-regex '^/api/(tags|correspondents|document_types|documents)/' --exclude-method POST --exclude-method PUT --exclude-method PATCH --exclude-method DELETE --phases examples,fuzzing\n").bold = True
    p_v1.add_run("• Lệnh chạy qua kịch bản chuẩn hóa:\n")
    r1_code = doc.add_paragraph()
    r1_code.add_run(".\\.venv\\Scripts\\python.exe scripts/run_round.py --round 1 --seed 404 --prefix TV4").font.name = "Consolas"

    # Vòng 2
    doc.add_heading("2.2. Vòng 2: Kiểm thử ghi dữ liệu (CRUD) trên Metadata — Chuẩn hóa 18 Operations", level=2).runs[0].font.name = "Times New Roman"
    p_v2 = doc.add_paragraph()
    p_v2.add_run("• Mục tiêu: Kiểm thử các thao tác tạo mới (POST), cập nhật (PUT/PATCH), xóa (DELETE) trên 3 nhóm metadata: tags, correspondents, document_types.\n• Cờ tham số: ").font.name = "Times New Roman"
    p_v2.add_run("--include-path-regex '^/api/(tags|correspondents|document_types)/' --checks all\n").bold = True
    p_v2.add_run("• Lệnh chạy:\n")
    r2_code = doc.add_paragraph()
    r2_code.add_run(".\\.venv\\Scripts\\python.exe scripts/run_round.py --round 2 --seed 404 --prefix TV4").font.name = "Consolas"

    # Vòng 3
    doc.add_heading("2.3. Vòng 3: Kiểm thử toàn diện mở rộng (Tất cả Checks) — Chuẩn hóa 53 Operations", level=2).runs[0].font.name = "Times New Roman"
    p_v3 = doc.add_paragraph()
    p_v3.add_run("• Mục tiêu: Bao gồm cả tài nguyên documents phức tạp và kích hoạt toàn bộ cơ chế kiểm tra lỗi nâng cao (--checks all).\n• Cờ tham số: ").font.name = "Times New Roman"
    p_v3.add_run("--include-path-regex '^/api/(tags|correspondents|document_types|documents)/' --checks all\n").bold = True
    p_v3.add_run("• Lệnh chạy:\n")
    r3_code = doc.add_paragraph()
    r3_code.add_run(".\\.venv\\Scripts\\python.exe scripts/run_round.py --round 3 --seed 404 --prefix TV4").font.name = "Consolas"

    doc.add_page_break()

    # MỤC 3: BẢNG SỐ LIỆU CÁC VÒNG CHẠY (MỤC 8.1 VÀ 8.2)
    h3 = doc.add_heading("3. Bảng số liệu đối chiếu 5 thành viên và Giải trình chuẩn hóa", level=1)
    h3.runs[0].font.name = "Times New Roman"
    h3.runs[0].font.color.rgb = RGBColor(24, 76, 120)

    doc.add_heading("3.1. Bảng số liệu các vòng chạy (Mục 8.1 Báo cáo nhóm)", level=2).runs[0].font.name = "Times New Roman"
    
    stat_table = doc.add_table(rows=6, cols=6)
    stat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Thành viên", "Seed", "Vòng 1 (19 op)", "Vòng 2 (18 op)", "Vòng 3 (53 op)", "Thời gian (V1 / V2 / V3)"]
    for j, h in enumerate(headers):
        cell = stat_table.cell(0, j)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.name = "Times New Roman"
        set_cell_background(cell, "2B4C7E")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)

    rows_data = [
        ("TV1", "101", "19 op, 99 test, 16 thất bại", "18 op, 1863 test, 52 thất bại", "53 op, 2016 test, 108 thất bại", "9,41 s / 194,09 s / 189,73 s"),
        ("TV2", "202", "[CHẠY LẠI] 19 op, 159 test, 14 thất bại\n(lần đầu: 65 op, 2517 test, 97 thất bại)", "18 op, 1481 test, 49 thất bại", "53 op, 2056 test, 108 thất bại", "31,85 s / 203,24 s / 211,21 s"),
        ("TV3", "303", "19 op, 74 test, 17 thất bại", "18 op, 744 test, 44 thất bại", "53 op, 1734 test, 108 thất bại", "2,71 s / 42,12 s / 102,41 s"),
        ("TV4 (Minh)", "404", "[CHẠY LẠI] 19 op, 72 test, 17 thất bại\n(lần đầu: 65 op, 8213 test, 111 thất bại)", "18 op, 1009 test, 47 thất bại", "53 op, 1516 test, 107 thất bại", "1,91 s / 49,71 s / 73,56 s\n(Process: 3,50s / 51,03s / 75,80s)"),
        ("TV5", "505", "19 op, 112 test, 16 thất bại", "18 op, 1940 test, 51 thất bại", "53 op, 1517 test, 104 thất bại", "11,65 s / 158,26 s / 132,31 s"),
    ]

    for i, row in enumerate(rows_data):
        for j, val in enumerate(row):
            cell = stat_table.cell(i+1, j)
            cell.text = val
            cell.paragraphs[0].runs[0].font.name = "Times New Roman"
            if "TV4" in row[0]:
                set_cell_background(cell, "EBF3FB")
                cell.paragraphs[0].runs[0].font.bold = True
            else:
                set_cell_background(cell, "F9F9F9" if i % 2 == 0 else "FFFFFF")
            set_cell_margins(cell)

    doc.add_paragraph()

    doc.add_heading("3.2. Nhận xét về tính nhất quán số liệu và sự cố 65 operations (Mục 8.2)", level=2).runs[0].font.name = "Times New Roman"
    p_nx = doc.add_paragraph()
    p_nx.add_run(
        "1. Giải thích số liệu 65 operations ban đầu của TV2 và TV4:\n"
        "   Lần chạy đầu tiên do cả TV2 và TV4 chỉ sử dụng cờ lọc phương thức `--include-method GET` mà chưa truyền `--include-path-regex`, "
        "dẫn đến việc Schemathesis quét toàn bộ 65 endpoint GET của cả hệ thống (vượt ra ngoài phạm vi 4 nhóm tài nguyên trọng tâm). "
        "Số liệu này đã được lưu giữ lại làm minh chứng khảo sát mở rộng (results/TV4_round1_seed404_thamdo.*).\n\n"
        "2. Kết quả sau chuẩn hóa của TV4:\n"
        "   Sau khi chạy lại với cờ regex chuẩn hóa `^/api/(tags|correspondents|document_types|documents)/`, số operation của TV4 đạt chuẩn chính xác 19 operations, "
        "sinh ra 72 test cases và phát hiện 17 test cases thất bại (22 unique failures) trong 1.91 giây (tổng process 3.50s). "
        "Số liệu này hoàn toàn nhất quán với TV3 (19 op, 74 test, 17 thất bại), TV1 (19 op, 99 test, 16 thất bại), TV5 (19 op, 112 test, 16 thất bại) và TV2 chạy lại (19 op, 159 test, 14 thất bại).\n\n"
        "3. Tính ổn định và khả năng tái lập (Reproducibility) ở Vòng 2 và Vòng 3:\n"
        "   Ở Vòng 2, cả 5 thành viên đều đạt đúng 18 operations với số lỗi thất bại dao động cực kỳ hẹp từ 44 đến 52 lỗi. "
        "Ở Vòng 3, cả 5 thành viên đều đạt đúng 53 operations với số lỗi thất bại ổn định từ 104 đến 108 lỗi. "
        "Điều này khẳng định kết quả kiểm thử hoàn toàn độc lập với hạt giống seed ngẫu nhiên và chứng minh các lỗi phát hiện là khiếm khuyết nội tại của Paperless-ngx v3.1.3."
    ).font.name = "Times New Roman"

    doc.add_page_break()

    # MỤC 4: DANH SÁCH BUGS VÀ NGUYÊN NHÂN GỐC
    h4 = doc.add_heading("4. Danh mục lỗi hệ thống (BUG-01 đến BUG-06) và Nguyên nhân gốc", level=1)
    h4.runs[0].font.name = "Times New Roman"
    h4.runs[0].font.color.rgb = RGBColor(24, 76, 120)

    bug_table = doc.add_table(rows=7, cols=5)
    bug_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_headers = ["Mã BUG", "Endpoint bị ảnh hưởng", "Loại lỗi kỹ thuật", "Nguyên nhân gốc (Root Cause)", "Tỷ lệ tái hiện"]
    for j, h in enumerate(b_headers):
        cell = bug_table.cell(0, j)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.name = "Times New Roman"
        set_cell_background(cell, "2B4C7E")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)

    bugs_data = [
        ("BUG-01", "POST /api/tags, correspondents, document_types", "Server Error (HTTP 500)", "serialisers.py: SetPermissionsMixin.validate_set_permissions() không kiểm tra set_permissions[action] có phải là dict hay không. Khi nhận boolean/integer gây TypeError sập máy chủ.", "4/5 TV (TV2, TV3, TV4, TV5)"),
        ("BUG-02", "GET /api/documents/?modified__year=...", "Server Error (HTTP 500)", "filters.py: DocumentFilterSet không đặt min/max cho modified__year. Giá trị <= 0 truyền vào hàm datetime(0, 1, 1) gây unhandled ValueError tại operations.py:636.", "3/5 TV (TV2, TV4, TV5)"),
        ("BUG-03", "POST /api/correspondents (và tags)", "Schema vi phạm hành vi", "views.py: get_queryset() bổ sung document_count qua .annotate(), nhưng phương thức tạo mới create() không đi qua get_queryset() nên response 201 Created thiếu trường bắt buộc document_count.", "4/5 TV (TV2, TV3, TV4, TV5)"),
        ("BUG-04", "GET /api/documents/{id}/share_links, notes", "Server Error (HTTP 500)", "views.py: Document.objects.get(pk=...) chỉ bắt DoesNotExist mà không bẫy ValueError khi {id} nhận chuỗi ký tự unicode hoặc chuỗi không phải số nguyên.", "4/5 TV (TV1, TV2, TV4, TV5)"),
        ("BUG-05", "Cả 4 nhóm endpoint API", "Schema thiếu sót", "OpenAPI Schema thiếu định nghĩa các mã lỗi thực tế (400, 404) và các Content-Type trả về thực tế như text/plain, image/webp tại download/preview.", "5/5 TV (Toàn nhóm)"),
        ("BUG-06", "GET /api/correspondents, document_types", "Schema vi phạm hành vi", "Response mã 200 thiếu trường permissions, đồng thời trường last_correspondence trả về giá trị null trái với định nghĩa bắt buộc kiểu chuỗi date trong Schema.", "4/5 TV (TV1, TV2, TV4, TV5)"),
    ]

    for i, row in enumerate(bugs_data):
        for j, val in enumerate(row):
            cell = bug_table.cell(i+1, j)
            cell.text = val
            cell.paragraphs[0].runs[0].font.name = "Times New Roman"
            set_cell_background(cell, "F9F9F9" if i % 2 == 0 else "FFFFFF")
            set_cell_margins(cell)

    doc.add_paragraph()

    # MỤC 5: HƯỚNG DẪN CHẠY KIỂM THỬ THỦ CÔNG 18 TEST CASES CỦA TV4
    h5 = doc.add_heading("5. Hướng dẫn chạy kiểm thử thủ công tự động (18 Test Cases Correspondents)", level=1)
    h5.runs[0].font.name = "Times New Roman"
    h5.runs[0].font.color.rgb = RGBColor(24, 76, 120)

    p_tc = doc.add_paragraph()
    p_tc.add_run(
        "Thành viên TV4 được phân công thiết kế chi tiết 18 ca kiểm thử hộp đen cho tài nguyên Correspondents (TC-CORR-01 đến TC-CORR-18). "
        "Toàn bộ 18 ca kiểm thử này đã được lập trình thành kịch bản Python tự động hoàn chỉnh tại `scripts/run_manual_tests.py`.\n\n"
        "Lệnh thực thi kiểm thử thủ công tự động:\n"
    ).font.name = "Times New Roman"

    tc_code = doc.add_paragraph()
    tc_code.add_run(".\\.venv\\Scripts\\python.exe scripts/run_manual_tests.py").font.name = "Consolas"

    p_tc_res = doc.add_paragraph()
    p_tc_res.add_run(
        "Kết quả thực thi:\n"
        "• 16/18 Test Cases ĐẠT (PASS) — Kiểm tra chặt chẽ các điều kiện biên, xác thực dữ liệu, bảo mật token, và lọc tìm kiếm.\n"
        "• 2/18 Test Cases BẮT ĐƯỢC LỖI THẬT (FAIL hợp lệ):\n"
        "  - TC-CORR-11: Gửi set_permissions sai kiểu làm máy chủ sập HTTP 500 (Tái hiện BUG-01 thành công).\n"
        "  - TC-CORR-01: Tạo mới thành công trả về 201 Created nhưng thiếu trường document_count (Tái hiện BUG-03 thành công)."
    ).font.name = "Times New Roman"

    # Lưu tài liệu
    output_path = "docs/HUONG_DAN_SU_DUNG_VA_KIEM_THU.docx"
    doc.save(output_path)
    print(f"[OK] Đã tạo thành công tài liệu Word: {output_path}")

if __name__ == "__main__":
    create_document()
